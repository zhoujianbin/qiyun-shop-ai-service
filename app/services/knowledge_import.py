"""知识文档解析与导入任务。

- 解析：front matter + 正文（MD/TXT/DOCX）；
- 编排：批量导入（sha256 去重、幂等键复用、单文件失败不阻塞）、切片预览、任务查询；
- 只做"文件 -> 文档/版本/切片"，向量化由 worker 异步补齐。
"""
from __future__ import annotations

import hashlib
import io
import pathlib
from dataclasses import dataclass, field
from datetime import date

import yaml
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.errors import NotFound, ValidationError
from app.integrations import storage
from app.models.generated import Chunk, Document, DocumentVersion, ImportJob
from app.services import audit
from app.services.knowledge_chunk import split_markdown

SUPPORTED = {".md", ".txt", ".docx"}
REQUIRED_META = ("document_id", "title", "category", "version", "effective_date", "owner", "status")
ALLOWED_STATUS = {"DRAFT", "REVIEWING", "ACTIVE", "ARCHIVED"}


class ParseError(ValueError):
    """单文件解析失败；message 会原样进入导入任务的 errors 清单。"""

    def __init__(self, message: str, *, line: int | None = None) -> None:
        super().__init__(message if line is None else f"第 {line} 行：{message}")
        self.line = line


@dataclass
class ParsedDocument:
    file_name: str
    meta: dict
    body: str
    first_body_line: int = 1
    warnings: list[str] = field(default_factory=list)

    @property
    def document_id(self) -> str:
        return str(self.meta["document_id"]).strip()

    @property
    def title(self) -> str:
        return str(self.meta["title"]).strip()

    @property
    def category(self) -> str:
        return str(self.meta["category"]).strip()

    @property
    def owner(self) -> str:
        return str(self.meta["owner"]).strip()

    @property
    def version(self) -> str:
        return str(self.meta["version"]).strip()

    @property
    def status(self) -> str:
        return str(self.meta["status"]).strip().upper()


def parse_document(file_name: str, payload: bytes) -> ParsedDocument:
    suffix = pathlib.Path(file_name).suffix.lower()
    if suffix not in SUPPORTED:
        raise ParseError(f"不支持的格式 {suffix or '未知'}，允许：{', '.join(sorted(SUPPORTED))}")
    if not payload:
        raise ParseError("文件内容为空")

    if suffix == ".docx":
        text = _docx_to_text(payload)
    else:
        try:
            text = payload.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise ParseError("文件不是 UTF-8 编码") from exc

    meta, body, first_line = _split_front_matter(text)
    _validate_meta(meta)
    if not body.strip():
        raise ParseError("正文为空")
    return ParsedDocument(file_name=file_name, meta=meta, body=body, first_body_line=first_line)


def _split_front_matter(text: str) -> tuple[dict, str, int]:
    lines = text.splitlines(keepends=True)
    if not lines or lines[0].strip() != "---":
        raise ParseError("缺少 front matter（文件必须以 --- 开头的元数据块起始）", line=1)
    closing = None
    for index in range(1, len(lines)):
        if lines[index].strip() == "---":
            closing = index
            break
    if closing is None:
        raise ParseError("front matter 未闭合（缺少结束的 ---）", line=1)
    try:
        meta = yaml.safe_load("".join(lines[1:closing])) or {}
    except yaml.YAMLError as exc:
        raise ParseError(f"front matter 不是合法 YAML：{exc}") from exc
    if not isinstance(meta, dict):
        raise ParseError("front matter 必须是键值对")
    body = "".join(lines[closing + 1:])
    return meta, body, closing + 2  # 正文第一行的绝对行号（1 基）


def _validate_meta(meta: dict) -> None:
    missing = [key for key in REQUIRED_META if not str(meta.get(key, "")).strip()]
    if missing:
        raise ParseError("front matter 缺少必填字段：" + "、".join(missing))
    status = str(meta["status"]).strip().upper()
    if status not in ALLOWED_STATUS:
        raise ParseError(f"front matter 的 status 非法：{status}（允许 {', '.join(sorted(ALLOWED_STATUS))}）")


def _docx_to_text(payload: bytes) -> str:
    """DOCX -> Markdown 近似文本：段落按行，表格转「表头: 值」。

    已知限制：语料中没有真实 DOCX，此解析器由 fixture 验证。
    DOCX 无 front matter，若正文未以元数据块开头，会在 _split_front_matter 处判为解析失败。
    """
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - 依赖缺失
        raise ParseError("未安装 python-docx，无法解析 DOCX") from exc
    try:
        document = Document(io.BytesIO(payload))
    except Exception as exc:  # noqa: BLE001 - 损坏文件统一转解析失败
        raise ParseError(f"DOCX 解析失败：{exc}") from exc

    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").lower()
        if style.startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "1"
            lines.append("#" * min(int(level), 6) + " " + text)
        else:
            lines.append(text)

    for table in document.tables:
        if not table.rows:
            continue
        header = [cell.text.strip() for cell in table.rows[0].cells]
        lines.append("")
        for row in table.rows[1:]:
            cells = [cell.text.strip() for cell in row.cells]
            pairs = [f"{h}: {v}" for h, v in zip(header, cells, strict=False) if h or v]
            if pairs:
                lines.append("；".join(pairs))
    return "\n".join(lines) + "\n"

# ---------------------------------------------------------------------------
# 导入任务编排
# ---------------------------------------------------------------------------
LIMITS = {"document_id": 48, "title": 200, "category": 48, "owner": 64}


def _sha256(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def _as_date(value):
    text = str(value or "").strip()
    if not text:
        return None
    try:
        return date.fromisoformat(text[:10])
    except ValueError:
        return None


def _check_lengths(parsed: ParsedDocument) -> None:
    for field, limit in LIMITS.items():
        value = str(parsed.meta.get(field, "")).strip()
        if len(value) > limit:
            raise ParseError(f"front matter 的 {field} 超过 {limit} 字符")


def preview_chunks(body: str, *, first_line: int = 1) -> list[dict]:
    """编辑器切片预览：只切不落库。"""
    return [
        {
            "ord": chunk.ord,
            "title_path": chunk.title_path,
            "content": chunk.content,
            "char_count": chunk.char_count,
            "line_start": chunk.line_start,
            "line_end": chunk.line_end,
        }
        for chunk in split_markdown(body, first_line=first_line)
    ]


def job_view(job: ImportJob) -> dict:
    return {
        "id": job.id,
        "source_type": job.source_type,
        "status": job.status,
        "total": job.total,
        "succeeded": job.succeeded,
        "failed": job.failed,
        "skipped": job.skipped,
        "errors": job.errors or [],
        "created_at": job.created_at.isoformat() if job.created_at else None,
    }


def create_job(
    db: Session,
    actor: dict,
    files: list[tuple[str, str | None, bytes]],
    *,
    idempotency_key: str | None = None,
    source_type: str = "MD_BUNDLE",
) -> dict:
    """批量导入：一次请求多文件，按内容去重、按幂等键复用任务、单文件失败不阻塞。"""
    if not files:
        raise ValidationError("没有上传任何文件")
    if source_type not in ("MD_BUNDLE", "RULE_YAML"):
        raise ValidationError(f"非法 source_type：{source_type}")

    # 上传层校验（类型 / 大小 / 空文件 / magic）失败 -> 整批 422
    for file_name, content_type, payload in files:
        storage.validate(file_name, content_type, payload)

    tenant_id = settings.tenant_id
    if idempotency_key:
        existing = db.execute(
            select(ImportJob).where(
                ImportJob.tenant_id == tenant_id, ImportJob.idempotency_key == idempotency_key
            )
        ).scalar_one_or_none()
        if existing is not None:
            return job_view(existing)

    job = ImportJob(
        tenant_id=tenant_id,
        source_type=source_type,
        status="RUNNING",
        total=len(files),
        errors=[],
        idempotency_key=idempotency_key,
        created_by=actor.get("id"),
    )
    db.add(job)
    db.flush()

    errors: list[dict] = []
    succeeded = failed = skipped = 0

    for file_name, _content_type, payload in files:
        try:
            parsed = parse_document(file_name, payload)
            _check_lengths(parsed)
        except ParseError as exc:
            failed += 1
            errors.append({"file": file_name, "reason": str(exc), "line": exc.line})
            continue

        checksum = _sha256(payload)
        document = db.execute(
            select(Document).where(
                Document.tenant_id == tenant_id, Document.document_id == parsed.document_id
            )
        ).scalar_one_or_none()

        if document is not None:
            duplicate = db.execute(
                select(DocumentVersion).where(
                    DocumentVersion.document_pk == document.id, DocumentVersion.checksum == checksum
                )
            ).scalar_one_or_none()
            if duplicate is not None:
                skipped += 1
                continue
        else:
            document = Document(
                tenant_id=tenant_id,
                document_id=parsed.document_id,
                title=parsed.title,
                category=parsed.category,
                owner=parsed.owner,
                status="DRAFT",
                channel_scope=[],
                identity_level="public",
                effective_date=_as_date(parsed.meta.get("effective_date")),
                metadata_json={"source_file": file_name},
            )
            db.add(document)
            db.flush()

        current = db.execute(
            select(func.max(DocumentVersion.version)).where(DocumentVersion.document_pk == document.id)
        ).scalar()
        version = DocumentVersion(
            tenant_id=tenant_id,
            document_pk=document.id,
            version=int(current or 0) + 1,
            content_md=parsed.body,
            checksum=checksum,
            status="DRAFT",
            source_kind="MD_BUNDLE",
            created_by=actor.get("id"),
        )
        db.add(version)
        db.flush()

        for chunk in split_markdown(parsed.body, first_line=parsed.first_body_line):
            db.add(Chunk(
                tenant_id=tenant_id,
                document_version_id=version.id,
                ord=chunk.ord,
                title_path=chunk.title_path or None,
                content=chunk.content,
                token_count=chunk.token_count,
                embedding_status="PENDING",
                identity_level=document.identity_level,
                metadata_json={
                    "source_line_start": chunk.line_start,
                    "source_line_end": chunk.line_end,
                    "headings": chunk.title_path,
                    "doc_version": str(parsed.version),
                    "source_file": file_name,
                },
            ))
        succeeded += 1

    job.succeeded = succeeded
    job.failed = failed
    job.skipped = skipped
    job.errors = errors
    if failed == 0:
        job.status = "DONE"
    elif succeeded == 0 and skipped == 0:
        job.status = "FAILED"
    else:
        job.status = "PARTIAL"

    audit.record(
        db, actor, "knowledge.import", "knowledge.import_job", job.id,
        after={"total": job.total, "succeeded": succeeded, "failed": failed, "skipped": skipped},
    )
    db.flush()
    return job_view(job)


def get_job(db: Session, job_id: int) -> dict:
    job = db.get(ImportJob, job_id)
    if job is None or job.tenant_id != settings.tenant_id:
        raise NotFound("导入任务不存在")
    return job_view(job)


def requeue_failed_chunks(db: Session) -> int:
    """把失败的向量化任务放回队列（供 worker --retry-failed 与运维脚本使用）。"""
    rows = db.execute(select(Chunk).where(Chunk.embedding_status == "FAILED")).scalars().all()
    for chunk in rows:
        chunk.embedding_status = "PENDING"
    db.flush()
    return len(rows)
